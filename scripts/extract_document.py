#!/usr/bin/env python3
"""
scripts/extract_document.py

Extrator multi-formato de alta fidelidade para ingestão de documentos em pipelines RAG.
Suporta:
  - PDF (.pdf)
  - Word (.docx, .doc)
  - PowerPoint (.pptx, .ppt)

Extrai estrutura hierárquica, páginas/slides delimitados, tabelas em formato Markdown,
e notas do apresentador (speaker notes).

Uso:
  python extract_document.py <caminho_documento> <arquivo_saida_txt> [--json-meta <meta.json>]
"""

import sys
import os
import json
import argparse
from datetime import datetime, timezone

def format_md_table(headers, rows):
    """Converte listas de células em uma tabela Markdown limpa."""
    if not headers and not rows:
        return ""
    if not headers and rows:
        headers = [f"Col {i+1}" for i in range(len(rows[0]))]
    
    clean_headers = [str(h).replace("\n", " ").replace("|", "\\|").strip() for h in headers]
    lines = ["| " + " | ".join(clean_headers) + " |"]
    lines.append("| " + " | ".join(["---"] * len(clean_headers)) + " |")
    
    for row in rows:
        clean_row = []
        for i in range(len(clean_headers)):
            cell = str(row[i]) if i < len(row) else ""
            clean_row.append(cell.replace("\n", " ").replace("|", "\\|").strip())
        lines.append("| " + " | ".join(clean_row) + " |")
    
    return "\n" + "\n".join(lines) + "\n\n"

def extract_pdf(file_path):
    """Extrai texto e metadados de PDF usando pypdf com fallback para markitdown."""
    meta = {
        "format": "PDF",
        "total_units": 0,
        "unit_name": "páginas",
        "has_tables": False,
        "has_notes": False,
        "title": "",
        "author": "",
    }
    content_parts = []
    
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        meta["total_units"] = len(reader.pages)
        
        info = reader.metadata
        if info:
            meta["title"] = str(info.get("/Title") or "").strip()
            meta["author"] = str(info.get("/Author") or "").strip()
        
        has_any_text = False
        has_any_images = False
        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                has_any_text = True
                content_parts.append(f"\n--- [PÁGINA {idx} DE {meta['total_units']}] ---\n\n{text}")
            else:
                content_parts.append(f"\n--- [PÁGINA {idx} DE {meta['total_units']}] ---\n\n[Página em branco ou apenas elementos visuais/imagem]")
            if hasattr(page, "images") and len(page.images) > 0:
                has_any_images = True
        meta["is_blank"] = (not has_any_text and not has_any_images)
    except Exception as e:
        # Fallback para markitdown
        try:
            from markitdown import MarkItDown
            md = MarkItDown()
            result = md.convert(file_path)
            content_parts = [result.text_content]
            meta["total_units"] = 1
        except Exception as e2:
            raise RuntimeError(f"Falha ao extrair PDF ({e}; fallback markitdown: {e2})")

    # Se a extração via pypdf foi muito vazia, tenta enriquecer com markitdown
    raw_joined = "\n".join(content_parts)
    if len(raw_joined.strip()) < 100:
        try:
            from markitdown import MarkItDown
            md = MarkItDown()
            res = md.convert(file_path)
            if len(res.text_content.strip()) > len(raw_joined.strip()):
                content_parts = [res.text_content]
        except Exception:
            pass

    return content_parts, meta

def extract_docx(file_path):
    """Extrai texto hierárquico e tabelas de DOCX usando python-docx."""
    meta = {
        "format": "DOCX",
        "total_units": 0,
        "unit_name": "seções",
        "has_tables": False,
        "has_notes": False,
        "title": "",
        "author": "",
    }
    content_parts = []
    
    try:
        import docx
        doc = docx.Document(file_path)
        
        if doc.core_properties:
            meta["title"] = doc.core_properties.title or ""
            meta["author"] = doc.core_properties.author or ""

        # Itera sobre elementos de nível de bloco mantendo ordem de parágrafos e tabelas
        section_count = 0
        current_section = []
        
        for element in doc.element.body:
            tag = element.tag.split("}")[-1]
            if tag == "p":
                from docx.text.paragraph import Paragraph
                p = Paragraph(element, doc)
                text = p.text.strip()
                if not text:
                    continue
                
                style_name = (p.style.name or "").lower()
                if "heading 1" in style_name or "título 1" in style_name:
                    section_count += 1
                    current_section.append(f"\n# {text}\n")
                elif "heading 2" in style_name or "título 2" in style_name:
                    current_section.append(f"\n## {text}\n")
                elif "heading 3" in style_name or "título 3" in style_name:
                    current_section.append(f"\n### {text}\n")
                elif "title" in style_name:
                    current_section.append(f"\n# {text}\n")
                elif "subtitle" in style_name:
                    current_section.append(f"\n_{text}_\n")
                elif "list" in style_name:
                    current_section.append(f"* {text}")
                else:
                    current_section.append(f"{text}\n")
            elif tag == "tbl":
                from docx.table import Table
                tbl = Table(element, doc)
                if tbl.rows:
                    meta["has_tables"] = True
                    headers = [c.text.strip() for c in tbl.rows[0].cells]
                    data_rows = []
                    for r in tbl.rows[1:]:
                        data_rows.append([c.text.strip() for c in r.cells])
                    md_table = format_md_table(headers, data_rows)
                    current_section.append(md_table)
        
        meta["total_units"] = max(1, section_count)
        content_parts = ["\n".join(current_section)]
    except Exception as e:
        # Fallback markitdown
        try:
            from markitdown import MarkItDown
            md = MarkItDown()
            res = md.convert(file_path)
            content_parts = [res.text_content]
            meta["total_units"] = 1
        except Exception as e2:
            raise RuntimeError(f"Falha ao extrair DOCX ({e}; fallback: {e2})")

    return content_parts, meta

def extract_pptx(file_path):
    """Extrai slides, títulos, bullets, tabelas e notas do apresentador usando python-pptx."""
    meta = {
        "format": "PPTX",
        "total_units": 0,
        "unit_name": "slides",
        "has_tables": False,
        "has_notes": False,
        "title": "",
        "author": "",
    }
    content_parts = []
    
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        meta["total_units"] = len(prs.slides)
        
        if prs.core_properties:
            meta["title"] = prs.core_properties.title or ""
            meta["author"] = prs.core_properties.author or ""

        for idx, slide in enumerate(prs.slides, start=1):
            slide_lines = []
            title_text = ""
            
            # 1. Título do slide
            if slide.shapes.title and slide.shapes.title.text:
                title_text = slide.shapes.title.text.strip()
            
            slide_header = f"--- [SLIDE {idx} DE {meta['total_units']}: {title_text or 'Sem Título'}] ---"
            slide_lines.append(slide_header)
            slide_lines.append("")
            
            # 2. Conteúdo das formas de texto e tabelas
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            # Indenta se for bullet secundário
                            level = getattr(paragraph, "level", 0)
                            indent = "  " * level
                            slide_lines.append(f"{indent}* {text}")
                elif shape.has_table:
                    meta["has_tables"] = True
                    table = shape.table
                    headers = [c.text.strip() for c in table.rows[0].cells]
                    rows = []
                    for r in table.rows[1:]:
                        rows.append([c.text.strip() for c in r.cells])
                    slide_lines.append(format_md_table(headers, rows))
            
            # 3. Notas do Apresentador (Speaker Notes) — Fonte crucial de contexto
            notes_text = ""
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                pass
            
            if notes_text:
                meta["has_notes"] = True
                slide_lines.append("")
                slide_lines.append("📌 **[NOTAS DO APRESENTADOR / CONTEXTO ADICIONAL]:**")
                for n_line in notes_text.splitlines():
                    if n_line.strip():
                        slide_lines.append(f"> {n_line.strip()}")
                slide_lines.append("")
            
            content_parts.append("\n".join(slide_lines))
    except Exception as e:
        # Fallback markitdown
        try:
            from markitdown import MarkItDown
            md = MarkItDown()
            res = md.convert(file_path)
            content_parts = [res.text_content]
            meta["total_units"] = 1
        except Exception as e2:
            raise RuntimeError(f"Falha ao extrair PPTX ({e}; fallback: {e2})")

    return content_parts, meta

def extract_html(file_path):
    """Extrai conteúdo de arquivos HTML/HTM com limpeza de ruídos (scripts/estilos) e converte em Markdown."""
    meta = {
        "format": "HTML",
        "total_units": 1,
        "unit_name": "páginas",
        "has_tables": False,
        "has_notes": False,
        "title": "",
        "author": "",
    }
    content_parts = []
    
    raw_html = ""
    for enc in ("utf-8", "latin-1", "cp1252", "iso-8859-1"):
        try:
            with open(file_path, "r", encoding=enc, errors="replace") as f:
                raw_html = f.read()
            break
        except Exception:
            continue

    if not raw_html:
        return [""], meta

    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(raw_html, "html.parser")
        
        if soup.title and soup.title.string:
            meta["title"] = soup.title.string.strip()
        elif soup.find("h1"):
            meta["title"] = soup.find("h1").get_text(strip=True)
            
        author_tag = soup.find("meta", attrs={"name": "author"})
        if author_tag and author_tag.get("content"):
            meta["author"] = author_tag.get("content").strip()

        for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "header"]):
            tag.decompose()

        if soup.find("table"):
            meta["has_tables"] = True

        import markdownify
        md_text = markdownify.markdownify(str(soup), heading_style="ATX").strip()
        
        header = f"# {meta['title'] or os.path.basename(file_path)}\n\n" if meta["title"] and not md_text.startswith("# ") else ""
        content_parts = [header + md_text]
    except Exception as e:
        try:
            from markitdown import MarkItDown
            md = MarkItDown()
            res = md.convert(file_path)
            content_parts = [res.text_content]
        except Exception as e2:
            raise RuntimeError(f"Falha ao extrair HTML ({e}; fallback: {e2})")

    return content_parts, meta

def extract_xlsx(file_path):
    """Extrai abas e tabelas de planilhas Excel (.xlsx) convertendo em Markdown."""
    meta = {
        "format": "SPREADSHEET",
        "total_units": 0,
        "unit_name": "abas",
        "has_tables": True,
        "has_notes": False,
        "title": os.path.basename(file_path),
        "author": "",
    }
    content_parts = []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        meta["total_units"] = len(wb.sheetnames)
        
        for idx, sheetname in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheetname]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            non_empty_rows = [r for r in rows if any(c is not None and str(c).strip() != "" for c in r)]
            if not non_empty_rows:
                continue
                
            headers = [str(c) if c is not None else "" for c in non_empty_rows[0]]
            clean_headers = [h.replace("\n", " ").replace("|", "\\|").strip() or f"Col_{i+1}" for i, h in enumerate(headers)]
            
            sheet_lines = [f"## [Aba {idx}/{meta['total_units']}: {sheetname}]\n"]
            sheet_lines.append("| " + " | ".join(clean_headers) + " |")
            sheet_lines.append("| " + " | ".join(["---"] * len(clean_headers)) + " |")
            
            for r in non_empty_rows[1:]:
                clean_row = [str(c).replace("\n", " ").replace("|", "\\|").strip() if c is not None else "" for c in r]
                while len(clean_row) < len(clean_headers):
                    clean_row.append("")
                sheet_lines.append("| " + " | ".join(clean_row[:len(clean_headers)]) + " |")
                
            content_parts.append("\n".join(sheet_lines))
    except Exception as e:
        try:
            from markitdown import MarkItDown
            md = MarkItDown()
            res = md.convert(file_path)
            content_parts = [res.text_content]
            meta["total_units"] = 1
        except Exception as e2:
            raise RuntimeError(f"Falha ao extrair XLSX ({e}; fallback: {e2})")

    return content_parts, meta

def extract_csv(file_path):
    """Extrai arquivos CSV/TSV convertendo em tabela Markdown."""
    meta = {
        "format": "CSV",
        "total_units": 1,
        "unit_name": "tabelas",
        "has_tables": True,
        "has_notes": False,
        "title": os.path.basename(file_path),
        "author": "",
    }
    content_parts = []
    
    raw_text = ""
    for enc in ("utf-8", "latin-1", "cp1252", "iso-8859-1"):
        try:
            with open(file_path, "r", encoding=enc, errors="replace") as f:
                raw_text = f.read()
            break
        except Exception:
            continue

    import csv, io
    delimiter = ","
    sample = raw_text[:3000]
    if "\t" in sample and sample.count("\t") > sample.count(","):
        delimiter = "\t"
    elif ";" in sample and sample.count(";") > sample.count(","):
        delimiter = ";"

    reader = csv.reader(io.StringIO(raw_text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return [""], meta

    headers = rows[0]
    clean_headers = [str(h).replace("\n", " ").replace("|", "\\|").strip() or f"Col_{i+1}" for i, h in enumerate(headers)]
    
    lines = [f"# {os.path.basename(file_path)}\n"]
    lines.append("| " + " | ".join(clean_headers) + " |")
    lines.append("| " + " | ".join(["---"] * len(clean_headers)) + " |")
    for row in rows[1:]:
        clean_row = [str(cell).replace("\n", " ").replace("|", "\\|").strip() for cell in row]
        while len(clean_row) < len(clean_headers):
            clean_row.append("")
        lines.append("| " + " | ".join(clean_row[:len(clean_headers)]) + " |")

    content_parts.append("\n".join(lines))
    return content_parts, meta

def extract_text(file_path):
    """Extrai texto simples e markdown (.txt, .md, .markdown)."""
    meta = {
        "format": "TEXT",
        "total_units": 1,
        "unit_name": "documento",
        "has_tables": False,
        "has_notes": False,
        "title": os.path.basename(file_path),
        "author": "",
    }
    raw_text = ""
    for enc in ("utf-8", "latin-1", "cp1252", "iso-8859-1"):
        try:
            with open(file_path, "r", encoding=enc, errors="replace") as f:
                raw_text = f.read()
            break
        except Exception:
            continue

    if "| ---" in raw_text:
        meta["has_tables"] = True

    return [raw_text], meta

def extract_json(file_path):
    """Extrai e formata dados estruturados JSON / JSONL."""
    meta = {
        "format": "JSON",
        "total_units": 1,
        "unit_name": "estruturas",
        "has_tables": False,
        "has_notes": False,
        "title": os.path.basename(file_path),
        "author": "",
    }
    raw_text = ""
    for enc in ("utf-8", "latin-1"):
        try:
            with open(file_path, "r", encoding=enc, errors="replace") as f:
                raw_text = f.read()
            break
        except Exception:
            continue

    content = f"# Arquivo de Dados: {os.path.basename(file_path)}\n\n"
    try:
        data = json.loads(raw_text)
        formatted = json.dumps(data, indent=2, ensure_ascii=False)
        content += f"```json\n{formatted}\n```"
    except Exception:
        lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
        meta["total_units"] = len(lines)
        content += f"```json\n{raw_text}\n```"

    return [content], meta

def extract_fallback_markitdown(file_path, fmt_name):
    """Fallback universal para outros formatos de mercado (.doc, .ppt, .xls, .rtf, .odt, .ods, .odp, .xml) via markitdown."""
    meta = {
        "format": fmt_name,
        "total_units": 1,
        "unit_name": "seções",
        "has_tables": False,
        "has_notes": False,
        "title": os.path.basename(file_path),
        "author": "",
    }
    from markitdown import MarkItDown
    md = MarkItDown()
    res = md.convert(file_path)
    return [res.text_content], meta

def main():
    parser = argparse.ArgumentParser(description="Extrator de documentos multi-formato para RAG")
    parser.add_argument("input_path", help="Caminho do documento de entrada (.pdf, .docx, .html, .pptx, .xlsx, .csv, .txt, etc.)")
    parser.add_argument("output_path", help="Caminho do arquivo TXT de saída com o texto estruturado")
    parser.add_argument("--json-meta", help="Caminho para gravar metadados em JSON", default=None)
    args = parser.parse_args()

    input_path = os.path.abspath(args.input_path)
    if not os.path.exists(input_path):
        print(f"Erro: arquivo não encontrado: {input_path}", file=sys.stderr)
        sys.exit(1)

    ext = os.path.splitext(input_path)[1].lower()
    
    if ext == ".pdf":
        parts, meta = extract_pdf(input_path)
    elif ext == ".docx":
        parts, meta = extract_docx(input_path)
    elif ext == ".pptx":
        parts, meta = extract_pptx(input_path)
    elif ext in (".html", ".htm", ".xhtml"):
        parts, meta = extract_html(input_path)
    elif ext in (".xlsx",):
        parts, meta = extract_xlsx(input_path)
    elif ext in (".csv", ".tsv"):
        parts, meta = extract_csv(input_path)
    elif ext in (".txt", ".md", ".markdown"):
        parts, meta = extract_text(input_path)
    elif ext in (".json", ".jsonl"):
        parts, meta = extract_json(input_path)
    elif ext in (".doc", ".ppt", ".xls", ".rtf", ".odt", ".ods", ".odp", ".xml"):
        parts, meta = extract_fallback_markitdown(input_path, ext.replace(".", "").upper())
    else:
        # Tentativa genérica com markitdown
        try:
            parts, meta = extract_fallback_markitdown(input_path, ext.replace(".", "").upper())
        except Exception as e:
            print(f"Erro: extensão não suportada ({ext}): {e}", file=sys.stderr)
            sys.exit(2)

    full_text = "\n\n".join(parts).strip()
    
    # Atualiza estatísticas textuais
    words = full_text.split()
    meta["char_count"] = len(full_text)
    meta["word_count"] = len(words)
    meta["file_name"] = os.path.basename(input_path)
    meta["extracted_at"] = datetime.now(timezone.utc).isoformat()

    # Salva o arquivo de texto estruturado
    os.makedirs(os.path.dirname(os.path.abspath(args.output_path)), exist_ok=True)
    with open(args.output_path, "w", encoding="utf-8") as f:
        f.write(full_text)

    # Salva metadados se solicitado
    if args.json_meta:
        os.makedirs(os.path.dirname(os.path.abspath(args.json_meta)), exist_ok=True)
        with open(args.json_meta, "w", encoding="utf-8") as f:
            json.dump(meta, f, indent=2, ensure_ascii=False)

    print(json.dumps(meta, ensure_ascii=False))

if __name__ == "__main__":
    main()
